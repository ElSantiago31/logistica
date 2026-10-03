"""Parche: agrega dropdown de validacion (DataValidation) para la columna
ETAPA en la plantilla oficial de importacion de operadores.

Restringe la celda a los 4 valores validos (previa, avanzada, evento,
desmontaje) y evita typos que caerian silenciosamente en 'evento'.
"""
import io
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "app" / "services" / "operator_import.py"

src = TARGET.read_text(encoding="utf-8")
orig = src

# 1. Imports de openpyxl adicionales
old_imports = (
    "from openpyxl import Workbook, load_workbook\n"
    "from openpyxl.styles import Font, PatternFill, Alignment\n"
)
new_imports = (
    "from openpyxl import Workbook, load_workbook\n"
    "from openpyxl.styles import Font, PatternFill, Alignment\n"
    "from openpyxl.utils import get_column_letter\n"
    "from openpyxl.worksheet.datavalidation import DataValidation\n"
)
if old_imports not in src:
    raise SystemExit("ERROR: bloque de imports openpyxl no encontrado")
src = src.replace(old_imports, new_imports, 1)

# 2. Dropdown en build_template, justo antes de freeze_panes
anchor = "    # Congelar primera fila\n    ws.freeze_panes = \"A2\"\n"
dv_block = (
    "    # Dropdown de validacion para ETAPA: restringe la celda a los valores\n"
    "    # validos (openpyxl DataValidation) y evita typos ('montaje') que\n"
    "    # caerian silenciosamente en 'evento' con solo un warning post-import.\n"
    "    stage_idx = next(i for i, c in enumerate(EXPECTED_COLUMNS, 1) if c[\"key\"] == \"stage\")\n"
    "    stage_letter = get_column_letter(stage_idx)\n"
    "    dv_stage = DataValidation(\n"
    "        type=\"list\",\n"
    "        formula1=f'\"{\",\".join(EVENT_STAGES)}\"',\n"
    "        allow_blank=True,\n"
    "        showErrorMessage=True,\n"
    "        errorTitle=\"Etapa no valida\",\n"
    "        error=\"Valores validos: previa, avanzada, evento, desmontaje (vacio = evento).\",\n"
    "        showInputMessage=True,\n"
    "        promptTitle=\"Etapa del operador\",\n"
    "        prompt=\"Selecciona previa, avanzada, evento o desmontaje. Vacio = evento.\",\n"
    "    )\n"
    "    ws.add_data_validation(dv_stage)\n"
    "    dv_stage.add(f\"{stage_letter}2:{stage_letter}1001\")\n"
    "\n"
)
if anchor not in src:
    raise SystemExit("ERROR: ancla freeze_panes no encontrada")
if "DataValidation(" in src.split("def build_template")[1].split("def ")[0] if "def build_template" in src else False:
    print("Ya aplicado; sin cambios.")
else:
    src = src.replace(anchor, dv_block + anchor, 1)

if src != orig:
    TARGET.write_text(src, encoding="utf-8")
    print("OK: dropdown ETAPA agregado a build_template()")
else:
    print("Sin cambios (ya aplicado).")

# Verificacion rapida
import openpyxl  # noqa: E402
import sys  # noqa: E402

sys.path.insert(0, str(ROOT))
from app.services.operator_import import build_template  # noqa: E402

raw = build_template()
wb = openpyxl.load_workbook(io.BytesIO(raw))
ws = wb.active
dvs = list(ws.data_validations.dataValidation)
assert dvs, "SIN DataValidation tras el parche!"
print(f"VERIFICADO: {len(dvs)} validacion(es); sqref={dvs[0].sqref}, formula={dvs[0].formula1}")